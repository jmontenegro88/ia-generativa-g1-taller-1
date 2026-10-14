from __future__ import annotations

import hashlib
import json
from pathlib import Path
from typing import Any

from src.config import (
    DATA_DIR,
    RAG_CHUNK_OVERLAP,
    RAG_CHUNK_SIZE,
    RAG_EMBEDDING_MODEL,
    RAG_TOP_K,
    RAG_VECTOR_STORE_DIR,
)

COLLECTION_NAME = "ecomarket_knowledge"
MANIFEST_NAME = "index_manifest.json"
POLICY_FORMAT_PRIORITY = (".pdf", ".docx", ".json", ".md")


def _policy_paths(data_dir: Path) -> list[Path]:
    candidates = [
        path
        for suffix in POLICY_FORMAT_PRIORITY
        for path in data_dir.glob(f"politica_*{suffix}")
    ]
    policies_by_category: dict[str, list[Path]] = {}
    for path in candidates:
        category = path.stem.removeprefix("politica_")
        policies_by_category.setdefault(category, []).append(path)

    priorities = {suffix: index for index, suffix in enumerate(POLICY_FORMAT_PRIORITY)}
    return [
        sorted(
            policies,
            key=lambda path: (priorities[path.suffix.lower()], path.name),
        )[0]
        for _, policies in sorted(policies_by_category.items())
    ]


def _read_policy_content(policy_path: Path) -> str:
    suffix = policy_path.suffix.lower()
    if suffix == ".md":
        content = policy_path.read_text(encoding="utf-8")
    elif suffix == ".pdf":
        from pypdf import PdfReader

        content = "\n".join(
            page.extract_text() or "" for page in PdfReader(str(policy_path)).pages
        )
    elif suffix == ".docx":
        from docx import Document as WordDocument

        content = "\n".join(
            paragraph.text
            for paragraph in WordDocument(str(policy_path)).paragraphs
            if paragraph.text.strip()
        )
    elif suffix == ".json":
        data = json.loads(policy_path.read_text(encoding="utf-8"))
        content = data.get("contenido") if isinstance(data, dict) else None
        if not isinstance(content, str):
            raise ValueError(
                f"{policy_path} debe contener un campo de texto 'contenido'."
            )
    else:
        raise ValueError(f"Formato de política no soportado: {policy_path.suffix}")

    if not content.strip():
        raise ValueError(f"No se pudo extraer texto de la política {policy_path}.")
    return content


def load_knowledge_documents(data_dir: Path = DATA_DIR) -> list[Any]:
    """Load policies and catalog from supported formats; orders remain exact-look-up data."""
    from langchain_core.documents import Document

    products_path = data_dir / "productos.json"
    products = json.loads(products_path.read_text(encoding="utf-8"))
    if not isinstance(products, list):
        raise ValueError(f"{products_path} debe contener una lista de productos.")

    policy_paths = _policy_paths(data_dir)
    if not policy_paths:
        raise FileNotFoundError(
            f"No se encontraron documentos politica_* compatibles en {data_dir}."
        )

    documents = []
    for policy_path in policy_paths:
        category = policy_path.stem.removeprefix("politica_")
        documents.append(
            Document(
                page_content=_read_policy_content(policy_path),
                metadata={
                    "tipo": "politica",
                    "categoria": category,
                    "fuente": f"data/{policy_path.name}",
                },
            )
        )

    for product in products:
        if not isinstance(product, dict) or not product.get("sku") or not product.get("name"):
            raise ValueError(f"Registro de producto inválido en {products_path}.")
        documents.append(
            Document(
                page_content=json.dumps(product, ensure_ascii=False, sort_keys=True),
                metadata={
                    "tipo": "producto",
                    "fuente": "data/productos.json",
                    "sku": str(product["sku"]),
                    "nombre": str(product["name"]),
                },
            )
        )
    return documents


def format_knowledge_context(documents: list[Any]) -> str:
    sections = []
    for document in documents:
        source = document.metadata.get("fuente", "fuente no identificada")
        sku = document.metadata.get("sku")
        if sku:
            source = f"{source} (SKU {sku})"
        sections.append(f"[Fuente: {source}]\n{document.page_content}")
    return "\n\n".join(sections)


def knowledge_sources(documents: list[Any]) -> list[str]:
    sources = []
    for document in documents:
        source = document.metadata.get("fuente")
        sku = document.metadata.get("sku")
        if source:
            label = f"{source} (SKU {sku})" if sku else source
            if label not in sources:
                sources.append(label)
    return sources


class KnowledgeBase:
    """Persistent Chroma index populated from the current policy and catalog."""

    def __init__(
        self,
        data_dir: Path = DATA_DIR,
        persist_directory: Path = RAG_VECTOR_STORE_DIR,
        embedding_model: str = RAG_EMBEDDING_MODEL,
    ) -> None:
        self.data_dir = data_dir
        self.persist_directory = persist_directory
        self.embedding_model = embedding_model
        self._vector_store: Any = None
        self._indexed_signature: str | None = None

    def search(
        self,
        query: str,
        document_type: str | None = None,
        category: str | None = None,
        k: int | None = None,
    ) -> list[Any]:
        if not query.strip():
            return []
        if k is not None and k < 1:
            raise ValueError("k debe ser mayor que cero.")
        vector_store = self._ensure_index()
        filters = []
        if document_type:
            filters.append({"tipo": {"$eq": document_type}})
        if category:
            filters.append({"categoria": {"$eq": category.strip().lower()}})
        search_filter = (
            filters[0]
            if len(filters) == 1
            else {"$and": filters} if filters else None
        )
        return vector_store.similarity_search(
            query,
            k=k or RAG_TOP_K,
            filter=search_filter,
        )

    def _source_signature(self) -> str:
        digest = hashlib.sha256()
        digest.update(self.embedding_model.encode("utf-8"))
        digest.update(str(RAG_CHUNK_SIZE).encode("ascii"))
        digest.update(str(RAG_CHUNK_OVERLAP).encode("ascii"))
        source_paths = _policy_paths(self.data_dir)
        source_paths.append(self.data_dir / "productos.json")
        for path in source_paths:
            filename = path.name
            digest.update(filename.encode("utf-8"))
            digest.update(path.read_bytes())
        return digest.hexdigest()

    def _ensure_index(self) -> Any:
        from langchain_chroma import Chroma
        from langchain_huggingface import HuggingFaceEmbeddings
        from langchain_text_splitters import RecursiveCharacterTextSplitter

        signature = self._source_signature()
        if self._vector_store is None:
            self.persist_directory.mkdir(parents=True, exist_ok=True)
            embeddings = HuggingFaceEmbeddings(
                model_name=self.embedding_model,
                model_kwargs={"device": "cpu"},
                encode_kwargs={"normalize_embeddings": True},
            )
            self._vector_store = Chroma(
                collection_name=COLLECTION_NAME,
                embedding_function=embeddings,
                persist_directory=str(self.persist_directory),
            )
        if self._indexed_signature == signature:
            return self._vector_store

        manifest_path = self.persist_directory / MANIFEST_NAME
        indexed_signature = None
        if manifest_path.exists():
            manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
            indexed_signature = manifest.get("signature")

        document_store = self._vector_store.get()
        old_ids = document_store.get("ids", [])
        needs_reindex = indexed_signature != signature or not old_ids
        if self._indexed_signature is not None:
            needs_reindex = self._indexed_signature != signature or not old_ids

        if needs_reindex:
            if old_ids:
                self._vector_store.delete(ids=old_ids)

            splitter = RecursiveCharacterTextSplitter(
                chunk_size=RAG_CHUNK_SIZE,
                chunk_overlap=RAG_CHUNK_OVERLAP,
            )
            chunks = splitter.split_documents(load_knowledge_documents(self.data_dir))
            ids = [
                hashlib.sha256(
                    (
                        chunk.metadata["fuente"]
                        + chunk.metadata.get("sku", "")
                        + str(index)
                        + chunk.page_content
                    ).encode("utf-8")
                ).hexdigest()
                for index, chunk in enumerate(chunks)
            ]
            self._vector_store.add_documents(chunks, ids=ids)
            manifest_path.write_text(
                json.dumps({"signature": signature}, ensure_ascii=False, indent=2),
                encoding="utf-8",
            )
            self._indexed_signature = signature
        else:
            self._indexed_signature = signature

        return self._vector_store
